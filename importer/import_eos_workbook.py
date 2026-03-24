#!/usr/bin/env python3
"""
Import EOS workbook sheets into matching stg.<branch>_* tables.

Design goals:
- Provenance-first: always populate _source_file, _source_sheet, _row_num.
- Reusable for historical backfill: supports dry-run and optional truncate.
- Conservative table matching: only imports sheets that map to existing stg tables.
- ENTRY sheets are treated as reference/master-data inputs; they are staged for
    audit and downstream reference modeling, not direct transaction-line loading.
"""

from __future__ import annotations

import argparse
import datetime
import os
import re
from decimal import Decimal
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Sequence, Tuple

import openpyxl
import pg8000.dbapi


def norm_token(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.lower())


@dataclass
class TableInfo:
    table_name: str
    all_columns: List[str]
    data_columns: List[str]
    data_column_types: List[str]


def get_pg_conn():
    return pg8000.dbapi.connect(
        host=os.getenv("PGHOST", "localhost"),
        port=int(os.getenv("PGPORT", "55432")),
        database=os.getenv("PGDATABASE", "exceldb"),
        user=os.getenv("PGUSER", "excel"),
        password=os.getenv("PGPASSWORD", "excelpass"),
    )


def fetch_branch_tables(cur, branch_prefix: str) -> Dict[str, TableInfo]:
    cur.execute(
        """
        SELECT table_name
        FROM information_schema.tables
        WHERE table_schema = 'stg'
          AND table_name LIKE %s
        ORDER BY table_name
        """,
        (f"{branch_prefix}_%",),
    )
    table_names = [r[0] for r in cur.fetchall()]

    table_map: Dict[str, TableInfo] = {}
    for table_name in table_names:
        cur.execute(
            """
            SELECT column_name, data_type
            FROM information_schema.columns
            WHERE table_schema = 'stg'
              AND table_name = %s
            ORDER BY ordinal_position
            """,
            (table_name,),
        )
        col_rows = cur.fetchall()
        all_cols = [r[0] for r in col_rows]
        data_cols = [c for c in all_cols if c not in ("_source_file", "_source_sheet", "_row_num")]
        type_by_col = {r[0]: r[1] for r in col_rows}
        data_types = [type_by_col[c] for c in data_cols]
        table_map[table_name] = TableInfo(
            table_name=table_name,
            all_columns=all_cols,
            data_columns=data_cols,
            data_column_types=data_types,
        )
    return table_map


def coerce_value_for_pg(value: object, pg_data_type: str) -> object:
    if value is None:
        return None

    if isinstance(value, str) and value.strip() == "":
        return None

    dtype = (pg_data_type or "").lower()

    if dtype in ("smallint", "integer", "bigint"):
        if isinstance(value, bool):
            return int(value)
        if isinstance(value, (int, float, Decimal)):
            return int(value)
        if isinstance(value, str):
            txt = value.strip().replace(",", "")
            if re.fullmatch(r"[+-]?[0-9]+", txt):
                return int(txt)
        return None

    if dtype in ("numeric", "decimal", "real", "double precision"):
        if isinstance(value, bool):
            return Decimal(int(value))
        if isinstance(value, (int, float, Decimal)):
            return Decimal(str(value))
        if isinstance(value, str):
            txt = value.strip().replace(",", "")
            if re.fullmatch(r"[+-]?[0-9]+(\.[0-9]+)?", txt):
                return Decimal(txt)
        return None

    if dtype == "boolean":
        if isinstance(value, bool):
            return value
        txt = str(value).strip().lower()
        if txt in ("1", "true", "t", "yes", "y"):
            return True
        if txt in ("0", "false", "f", "no", "n"):
            return False
        return None

    if dtype == "date":
        if isinstance(value, datetime.datetime):
            return value.date()
        if isinstance(value, datetime.date):
            return value
        txt = str(value).strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
            try:
                return datetime.datetime.strptime(txt, fmt).date()
            except ValueError:
                continue
        return None

    if dtype.startswith("timestamp"):
        if isinstance(value, datetime.datetime):
            return value
        if isinstance(value, datetime.date):
            return datetime.datetime.combine(value, datetime.time())
        txt = str(value).strip()
        try:
            return datetime.datetime.fromisoformat(txt.replace("Z", "+00:00"))
        except ValueError:
            return None

    return str(value)


def match_sheet_to_table(branch_prefix: str, sheet_name: str, table_map: Dict[str, TableInfo]) -> str | None:
    token = norm_token(sheet_name)
    direct = f"{branch_prefix}_{token}"
    if direct in table_map:
        return direct

    # Secondary pass: match normalized suffix against existing table suffixes.
    candidates = []
    for table_name in table_map:
        suffix = table_name[len(branch_prefix) + 1 :]
        if norm_token(suffix) == token:
            candidates.append(table_name)
    if len(candidates) == 1:
        return candidates[0]
    return None


def row_has_data(values: Sequence[object]) -> bool:
    for v in values:
        if v is None:
            continue
        if isinstance(v, str) and v.strip() == "":
            continue
        return True
    return False


def is_numeric_like(value: object) -> bool:
    if value is None:
        return False
    if isinstance(value, (int, float)):
        return True
    if isinstance(value, str):
        txt = value.strip().replace(",", "")
        return re.fullmatch(r"[+-]?[0-9]+(\.[0-9]+)?", txt) is not None
    return False


def is_likely_pay_transaction_row(row_values_ao: Sequence[object]) -> bool:
    if len(row_values_ao) < 15:
        return False

    dept = row_values_ao[0]
    business_date = row_values_ao[1]
    ex_vat = row_values_ao[11]
    vat = row_values_ao[12]
    total = row_values_ao[13]

    dept_txt = dept.strip().lower() if isinstance(dept, str) else ""
    date_txt = business_date.strip().lower() if isinstance(business_date, str) else ""

    # Skip obvious header rows.
    if dept_txt in {"department", "dept"}:
        return False
    if date_txt in {"date", "paid date"}:
        return False

    has_valid_date = isinstance(business_date, (datetime.date, datetime.datetime))
    has_numeric_amount = is_numeric_like(ex_vat) or is_numeric_like(vat) or is_numeric_like(total)

    # Keep rows that look like transaction lines by date and/or amount signal.
    return has_valid_date or has_numeric_amount


def prepare_row_payload(
    row_values: Sequence[object],
    source_file: str,
    source_sheet: str,
    row_num: int,
    data_column_types: Sequence[str],
    data_column_count: int,
    map_by_position: bool,
    mapped_column_count_override: int | None,
) -> Tuple[object, ...]:
    if not map_by_position:
        return (source_file, source_sheet, row_num)

    target_count = data_column_count
    if mapped_column_count_override is not None:
        target_count = min(data_column_count, mapped_column_count_override)

    mapped = []
    for idx, raw_value in enumerate(row_values[:target_count]):
        mapped.append(coerce_value_for_pg(raw_value, data_column_types[idx]))

    if len(mapped) < target_count:
        mapped.extend([None] * (target_count - len(mapped)))

    # Pad to full table width so untouched columns remain NULL.
    if target_count < data_column_count:
        mapped.extend([None] * (data_column_count - target_count))

    return (source_file, source_sheet, row_num, *mapped)


def build_insert_sql(table_name: str, data_columns: Sequence[str], map_by_position: bool) -> str:
    cols = ["_source_file", "_source_sheet", "_row_num"]
    if map_by_position:
        cols.extend(data_columns)
    col_sql = ", ".join(f'"{c}"' for c in cols)
    values_sql = ", ".join(["%s"] * len(cols))
    return f"INSERT INTO stg.{table_name} ({col_sql}) VALUES ({values_sql})"


def truncate_table(cur, table_name: str) -> None:
    cur.execute(f"TRUNCATE TABLE stg.{table_name}")


def import_workbook(
    workbook_path: Path,
    branch_prefix: str,
    dry_run: bool,
    truncate_matched: bool,
    map_by_position: bool,
    pay_ao_only: bool,
) -> None:
    source_file = workbook_path.name

    conn = get_pg_conn()
    try:
        conn.autocommit = False
        cur = conn.cursor()
        table_map = fetch_branch_tables(cur, branch_prefix)

        if not table_map:
            raise RuntimeError(f"No staging tables found for prefix stg.{branch_prefix}_*")

        wb = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)

        matched: List[Tuple[str, str]] = []
        unmatched: List[str] = []
        inserted_counts: Dict[str, int] = {}

        for ws in wb.worksheets:
            table_name = match_sheet_to_table(branch_prefix, ws.title, table_map)
            if not table_name:
                unmatched.append(ws.title)
                continue

            tinfo = table_map[table_name]
            matched.append((ws.title, table_name))

            if truncate_matched and not dry_run:
                truncate_table(cur, table_name)

            payload_rows: List[Tuple[object, ...]] = []
            mapped_column_count_override: int | None = None
            effective_map_by_position = map_by_position
            # PAY transaction lines are A:O (15 columns). Ignore P:AC and AD:AJ summary region.
            if pay_ao_only and table_name.endswith("pay"):
                mapped_column_count_override = 15
            if pay_ao_only and not table_name.endswith("pay"):
                effective_map_by_position = False

            for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                row_scope = row
                if pay_ao_only and table_name.endswith("pay"):
                    row_scope = row[:15]

                if not row_has_data(row_scope):
                    continue

                if pay_ao_only and table_name.endswith("pay") and not is_likely_pay_transaction_row(row_scope):
                    continue

                payload_rows.append(
                    prepare_row_payload(
                        row_values=row,
                        source_file=source_file,
                        source_sheet=ws.title,
                        row_num=idx,
                        data_column_types=tinfo.data_column_types,
                        data_column_count=len(tinfo.data_columns),
                        map_by_position=effective_map_by_position,
                        mapped_column_count_override=mapped_column_count_override,
                    )
                )

            if dry_run:
                inserted_counts[table_name] = inserted_counts.get(table_name, 0) + len(payload_rows)
                continue

            if payload_rows:
                sql = build_insert_sql(
                    table_name,
                    tinfo.data_columns,
                    map_by_position=effective_map_by_position,
                )
                cur.executemany(sql, payload_rows)

            inserted_counts[table_name] = inserted_counts.get(table_name, 0) + len(payload_rows)

        if dry_run:
            conn.rollback()
        else:
            conn.commit()
    finally:
        conn.close()

    print("=== EOS Workbook Import Summary ===")
    print(f"Workbook: {workbook_path}")
    print(f"Branch prefix: {branch_prefix}")
    print(f"Dry run: {dry_run}")
    print(f"Truncate matched tables first: {truncate_matched}")
    print(f"Map non-source columns by position: {map_by_position}")
    print(f"PAY A:O only mapping mode: {pay_ao_only}")
    print("")
    print("Matched sheets -> tables:")
    for sheet_name, table_name in matched:
        print(f"  {sheet_name} -> stg.{table_name}")
    print("")
    print("Inserted rows per table (or would insert in dry-run):")
    for table_name in sorted(inserted_counts):
        print(f"  stg.{table_name}: {inserted_counts[table_name]}")
    print("")
    if unmatched:
        print("Unmatched sheets (no stg table mapping):")
        for sheet_name in unmatched:
            print(f"  {sheet_name}")
    else:
        print("Unmatched sheets: none")


def main() -> None:
    parser = argparse.ArgumentParser(description="Import EOS workbook into matching stg tables")
    parser.add_argument("--workbook", required=True, help="Path to EOS workbook (.xlsx)")
    parser.add_argument("--branch", default="eosz", choices=["eosz", "eosq", "eosblum"], help="Staging table branch prefix")
    parser.add_argument("--dry-run", action="store_true", help="Parse and map workbook but do not commit inserts")
    parser.add_argument(
        "--truncate-matched",
        action="store_true",
        help="Truncate matched stg.<branch>_* tables before inserting",
    )
    parser.add_argument(
        "--map-by-position",
        action="store_true",
        help="Also map worksheet row values into non-source table columns by ordinal position",
    )
    parser.add_argument(
        "--pay-ao-only",
        action="store_true",
        help="When mapping by position, map PAY sheets using only columns A:O (15 columns)",
    )

    args = parser.parse_args()
    workbook_path = Path(args.workbook)
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    import_workbook(
        workbook_path=workbook_path,
        branch_prefix=args.branch,
        dry_run=args.dry_run,
        truncate_matched=args.truncate_matched,
        map_by_position=args.map_by_position,
        pay_ao_only=args.pay_ao_only,
    )


if __name__ == "__main__":
    main()
